#!/usr/bin/env python3
"""Generate an ATS-readable CV as .docx.

Applicant tracking systems parse a document by walking its text in order. They
are defeated by the things that make a CV look designed:

  * tables and columns   - text is read in the wrong order, or not at all
  * text boxes           - frequently skipped entirely
  * headers and footers  - many parsers ignore them, so contact details put
                           there simply vanish
  * icons and images     - carry no text
  * unusual fonts        - can extract as mojibake

So this is deliberately a single column of plain paragraphs with standard
headings ("Work Experience", "Education", "Skills") that parsers match on, and
contact details in the body rather than a header.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.dirname(os.path.abspath(__file__))
ACCENT = RGBColor(0x1F, 0x4E, 0x33)


def setup(doc):
    for section in doc.sections:
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"      # a font every parser handles
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1.05


def para(doc, text="", size=10.5, bold=False, italic=False, color=None,
         space_before=0, space_after=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return p


def heading(doc, text):
    """A plain bold uppercase line. ATS matches section names literally, so
    the wording here is the conventional one rather than anything clever."""
    p = para(doc, text.upper(), size=11.5, bold=True, color=ACCENT,
             space_before=12, space_after=2)
    pb = p.paragraph_format
    pb.keep_with_next = True
    return p


def rule(doc):
    para(doc, "_" * 96, size=6, color=RGBColor(0xB8, 0xC4, 0xBC), space_after=6)


def role(doc, title, org, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{title}")
    r.bold = True
    r.font.size = Pt(10.5)
    r2 = p.add_run(f"  |  {org}  |  {dates}")
    r2.font.size = Pt(10.5)
    for b in bullets:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(0)
        bp.paragraph_format.left_indent = Inches(0.22)
        bp.add_run(b).font.size = Pt(10.5)


def build():
    doc = Document()
    setup(doc)

    # ---------- identity: in the body, never a header ----------
    para(doc, "OSCAR MURETI", size=20, bold=True, color=ACCENT, space_after=1)
    para(doc, "Software Developer — Mobile, Web, AI & Machine Learning",
         size=11.5, space_after=3)
    para(doc, "Nairobi, Kenya  |  +254 707 817 900  |  oscarmureti75@gmail.com",
         size=10, space_after=1)
    para(doc, "https://oscarmureti.github.io  |  https://github.com/oscarmureti",
         size=10, space_after=2)
    rule(doc)

    heading(doc, "Professional Summary")
    para(doc,
         "Software developer with seven years building and shipping production "
         "applications across mobile, web and machine learning. Published apps "
         "on Google Play and the App Store, including fintech, public-health and "
         "e-commerce systems used by county governments and businesses in Kenya. "
         "Comfortable owning a product end to end: architecture, implementation, "
         "automated testing, CI/CD and store release.", space_after=2)

    heading(doc, "Technical Skills")
    for label, items in [
        ("Mobile", "Flutter, Dart, Android, iOS, React Native"),
        ("Web", "React.js, JavaScript, PHP, Bootstrap, REST APIs, HTML/CSS"),
        ("AI & Machine Learning", "Python, machine learning, data analysis"),
        ("Data & Backend", "Firebase, MySQL, DHIS2, Java, REST API design"),
        ("Practices & Tooling",
         "Git, GitHub Actions, CI/CD, automated testing, Google Play Console, "
         "App Store Connect, Xcode, Android Studio, server management"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(f"{label}: ")
        r.bold = True
        r.font.size = Pt(10.5)
        p.add_run(items).font.size = Pt(10.5)

    heading(doc, "Work Experience")
    role(doc, "Mobile Developer (iOS and Android)", "Nobuk Africa",
         "May 2022 – Present", [
         "Build and maintain Android and iOS applications for a fintech platform "
         "providing loans to small-scale businesses.",
         "Work across the full mobile stack in Flutter, from feature "
         "implementation through release to both app stores.",
         ])
    role(doc, "Mobile Developer (iOS and Android)",
         "Bright Coders Factory East Africa", "January 2022 – Present", [
         "Develop Android and iOS applications for a fintech product, working "
         "within a distributed team.",
         ])
    role(doc, "CRO / Mobile Developer (iOS and Android)", "Cryosoft Corporation",
         "September 2018 – Present", [
         "Lead mobile development across the company's product portfolio, "
         "including Liquid, CoolPam, Ride Africa, Soi Pallets, Muvi and Music "
         "Family.",
         "Deliver both customer-facing and staff-facing applications for the "
         "same products, covering ordering, fulfilment and reporting.",
         "Built and published multiple applications to Google Play.",
         ])
    role(doc, "Software Developer", "Vihiga County Referral Hospital",
         "January 2021 – August 2021", [
         "Built a Flutter application for tracking diabetes and hypertension "
         "patients across Vihiga County under the HealthIT project.",
         "Enabled clinicians to monitor patient health indicators remotely and "
         "prompted patients to record readings and attend appointments, aimed at "
         "reducing complications from untracked indicators.",
         ])
    role(doc, "Software Developer", "Kisumu County Ministry of Health",
         "June 2019 – August 2019", [
         "Developed Kisumu Youth, a mobile platform connecting young people in "
         "the county with mentors and enabling them to share ideas and promote "
         "their businesses. Published to Google Play.",
         ])
    role(doc, "DHIS2 Applications Developer", "HI4Kenya Bootcamp", "2018", [
         "Contributed to Wajibika, a React Native data-collection tool used by "
         "the Ministry of Health, storing collected data into DHIS2.",
         ])

    heading(doc, "Selected Projects")
    role(doc, "Puzzle Game Portfolio (Wordroot, Solitaire Assured, Sudoku Mentor)",
         "Personal", "2026", [
         "Three published Flutter games on a shared engine, with monetisation, "
         "in-app purchases, localisation and progression built once and reused.",
         "Implemented a Klondike solver that proved all 4,023 shipped deals "
         "winnable before release, and a sudoku engine that grades 5,493 puzzles "
         "by the solving technique each requires rather than by clue count.",
         "Localised to 12 languages, covered by 44 automated tests, and released "
         "through a GitHub Actions pipeline that deploys to Google Play and "
         "TestFlight on merge.",
         ])
    role(doc, "DairyVibes", "Android, published", "", [
         "Dairy farm management application covering milk records, herd, health "
         "and feed tracking. Built with Flutter and Firebase.",
         ])
    role(doc, "CoolPam Management", "Cryosoft Corporation", "", [
         "Water management system spanning Android and iOS apps, an admin "
         "website in React.js and a desktop administration tool.",
         ])
    role(doc, "Liquid and Liquid Staff", "Cryosoft Corporation", "", [
         "Food-delivery marketplace plus the staff application handling order "
         "receipt, processing, delivery and revenue reporting. Flutter, "
         "Firebase, PHP and Google Maps.",
         ])

    heading(doc, "Education")
    role(doc, "BSc, Mathematics and Computer Science", "Maseno University",
         "2016 – 2020", [])

    heading(doc, "Languages")
    para(doc, "English (fluent)  |  Kiswahili (native)", space_after=2)

    path = os.path.join(OUT, "Oscar-Mureti-CV.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    p = build()
    print(f"  wrote {p} ({os.path.getsize(p)/1024:.0f} KB)")
